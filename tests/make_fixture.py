"""Builds FAKE Excel files that copy the layout of KBA's FZ 10.1 sheet.

The numbers are invented and exist only to test the parser. They are never
used for any analysis or chart in this repository.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import Workbook

MONTH_DE = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli",
            "August", "September", "Oktober", "November", "Dezember"]

# 14 groups, in the same order as the real file. Each has 3 columns.
GROUPS = [
    "Insgesamt", "mit Dieselantrieb", "mit Hybridantrieb\n(incl. Plug-in-Hybrid)",
    "Benzin-Hybridantrieb\n(incl. Plug-in-Hybrid)", "Diesel-Hybridantrieb\n(incl. Plug-in-Hybrid)",
    "Hybridantrieb\n(ohne Plug-in-Hybrid)", "Benzin-Hybridantrieb\n(ohne Plug-in-Hybrid)",
    "Diesel-Hybridantrieb\n(ohne Plug-in-Hybrid)", "Plug-in-Hybridantrieb",
    "Benzin-Plug-in-Hybridantrieb", "Diesel-Plug-in-Hybridantrieb",
    "mit Elektroantrieb (BEV)", "mit Allradantrieb", "Cabriolets",
]

BRANDS = {
    "ALFA ROMEO": ["GIULIA", "JUNIOR", "SONSTIGE"],
    "AUDI": ["A3", "A6", "Q4", "SONSTIGE"],
    "BMW": ["1ER", "3ER", "X1", "SONSTIGE"],
}


def _fmt(n: int):
    """The real file shows zero as '-' (text)."""
    return "-" if n == 0 else int(n)


def make_month_file(path: Path, year: int, month: int, seed: int = 0, bev_boost: float = 0.0):
    rng = np.random.default_rng(seed + year * 100 + month)
    wb = Workbook()
    wb.active.title = "Deckblatt"
    wb.create_sheet("Impressum")
    wb.create_sheet("Inhaltsverzeichnis")
    ws = wb.create_sheet("FZ 10.1")

    ws["B2"] = "zurück zum Inhaltsverzeichnis"
    ws["B4"] = "Fahrzeugzulassungen (FZ)"
    ws["B5"] = "Neuzulassungen von Personenkraftwagen nach Marken und Modellreihen"
    ws["B6"] = f"FZ 10.1 Neuzulassungen von Personenkraftwagen nach Marken und Modellreihen im {MONTH_DE[month-1]} {year}"

    ws.cell(8, 2, "Marke")
    ws.cell(8, 3, "Modellreihe")
    for g, name in enumerate(GROUPS):
        c = 4 + 3 * g
        ws.cell(8, c, name)
        ws.cell(9, c, f"{MONTH_DE[month-1]} {year}")
        ws.cell(9, c + 1, f"Jan. - {MONTH_DE[month-1]} {year}")
        ws.cell(9, c + 2, "Anteil in %")

    r = 10
    grand = np.zeros(len(GROUPS))
    for brand, models in BRANDS.items():
        block = np.zeros(len(GROUPS))
        for i, model in enumerate(models):
            total = int(rng.integers(200, 3000))
            bev = int(total * min(0.6, rng.uniform(0.02, 0.15) + bev_boost))
            phev = int(total * rng.uniform(0.02, 0.12))
            hev = int(total * rng.uniform(0.05, 0.25))
            diesel = int(total * rng.uniform(0.0, 0.10))
            row = np.zeros(len(GROUPS))
            row[0], row[1], row[5], row[8], row[11] = total, diesel, hev, phev, bev
            row[2] = hev + phev                      # hybrid incl. plug-in overlaps hev + phev
            row[3] = row[2] * 0.9
            row[4] = row[2] - row[3]
            row[12] = total * 0.3
            block += row
            if i == 0:
                ws.cell(r, 2, brand)                 # brand only on the first row of a block
            ws.cell(r, 3, model)
            for g in range(len(GROUPS)):
                ws.cell(r, 4 + 3 * g, _fmt(int(row[g])))
                ws.cell(r, 5 + 3 * g, _fmt(int(row[g] * 6)))
                ws.cell(r, 6 + 3 * g, 1.0)
            r += 1
        ws.cell(r, 2, f"{brand} ZUSAMMEN")           # subtotal row (must be ignored)
        for g in range(len(GROUPS)):
            ws.cell(r, 4 + 3 * g, _fmt(int(block[g])))
        grand += block
        r += 1
    # "other brands": only a subtotal line, no model rows (as in the real newer files)
    other = np.zeros(len(GROUPS))
    other[0], other[1], other[11] = 500, 20, 30
    ws.cell(r, 2, "SONSTIGE ZUSAMMEN")
    for g in range(len(GROUPS)):
        ws.cell(r, 4 + 3 * g, _fmt(int(other[g])))
    grand += other
    r += 1
    ws.cell(r, 2, "NEUZULASSUNGEN INSGESAMT")
    for g in range(len(GROUPS)):
        ws.cell(r, 4 + 3 * g, _fmt(int(grand[g])))
    ws.cell(r + 2, 2, "Quelle: Kraftfahrt-Bundesamt (Fußnote, muss ignoriert werden)")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return grand


def make_series(folder: Path, start=(2024, 9), months=24):
    """Creates `months` consecutive fake monthly files. BEV share drifts upward."""
    y, m = start
    for i in range(months):
        make_month_file(folder / f"fz10_{y}_{m:02d}.xlsx", y, m, bev_boost=0.008 * i)
        m += 1
        if m == 13:
            y, m = y + 1, 1
